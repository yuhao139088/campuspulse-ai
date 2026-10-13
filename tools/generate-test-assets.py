from pathlib import Path
import csv
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

destination = Path(sys.argv[1]).resolve()
destination.mkdir(parents=True, exist_ok=True)
headers = ["activity", "month", "registered", "checkedIn", "budget", "satisfaction"]
rows = [
    ["Welcome", "2026-05", 80, 68, 1800, 4.5],
    ["Workshop", "2026-05", 45, 40, 950, 4.7],
    ["OpenDay", "2026-06", 100, 83, 2500, 4.2],
    ["Workshop", "2026-06", 56, 49, 1100, 4.8],
    ["Welcome", "2026-07", 90, 76, 2000, 4.6],
    ["Workshop", "2026-07", 60, 52, 1250, 4.7],
    ["OpenDay", "2026-08", 120, 84, 3600, 3.9],
    ["Welcome", "2026-08", 95, 82, 2300, 4.6],
    ["Workshop", "2026-09", 70, 61, 1450, 4.8],
    ["Welcome", "2026-09", 110, 96, 2400, 4.7],
    ["OpenDay", "2026-10", 160, 60, 5800, 3.1],
    ["Workshop", "2026-10", 75, 66, 1550, 4.6],
]
def export(name, columns, values):
    with (destination / name).open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerows(values)

export("01-campus-operations.csv", headers, rows)
export("03-duplicates-missing-outliers.csv", headers, rows + [rows[0], ["Incomplete", "2026-10", 40, "", 700, ""], ["Outlier", "2026-10", 40, 5000, 99999, 1.2]])
export("04-duplicate-headers-rejected.csv", ["activity", "registered", "registered"], [["Example", 30, 25]])
export("05-chinese-columns.csv", ["活动", "月份", "报名人数", "签到人数", "预算", "满意度"], rows)
export("06-private-space-isolation.csv", ["marker", "value"], [["PRIVATE_TEST_ACCOUNT_B", 87123]])

book = Workbook()
sheet = book.active
sheet.title = "Operations"
sheet.append(headers)
for row in rows:
    sheet.append(row)
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
for column in range(1, len(headers) + 1):
    cell = sheet.cell(1, column)
    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="153449")
    cell.comment = Comment("Synthetic test fixture; these are not actual campus records. Edit input cells to test imports.", "CampusPulse AI")
    sheet.column_dimensions[cell.column_letter].width = 20
for row in sheet.iter_rows(min_row=2):
    for cell in row:
        cell.font = Font(name="Arial", color="153449")
        cell.alignment = Alignment(vertical="center")
instructions = book.create_sheet("Read me")
for text in ["Synthetic CampusPulse test data", "Import reads the first worksheet: Operations.", "All numeric values are generated test inputs, not real student data.", "Edit rows 2 onward to test changed data. Header names are required and unique.", "No formulas are used: this file is an input data fixture, not a calculated report."]:
    instructions.append([text])
    instructions.cell(instructions.max_row, 1).font = Font(name="Arial", size=11)
instructions.column_dimensions["A"].width = 105
book.save(destination / "02-campus-operations.xlsx")
print("Created 6 original spreadsheet fixtures; 12 clean rows; registered=" + str(sum(row[2] for row in rows)) + "; checkedIn=" + str(sum(row[3] for row in rows)))
